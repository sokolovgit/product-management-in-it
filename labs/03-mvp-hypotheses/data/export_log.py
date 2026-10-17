"""Log-файл Concierge MVP у xlsx: журнал подій, щоденник тестування і підсумки по репетиторах."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from analyze_log import load

HERE = Path(__file__).parent
OUT = HERE / "concierge_log.xlsx"

EVENTS = {
    "lesson": "заняття",
    "payment": "оплата",
    "reminder_first": "нагадування",
    "reminder_repeat": "повторне нагадування",
    "reminder_false": "нагадування (оплату вже внесено)",
    "reply_negative": "негативна відповідь",
    "reply_paid_already": "відповідь «вже оплатили»",
    "tutor_report": "звіт репетитора",
    "weekly_summary": "підсумок тижня",
}


def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(
            10, max(len(str(c.value or "")) for c in col) + 2
        )


def main():
    rows = load()
    m = json.loads((HERE / "metrics.json").read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    sheet(
        wb,
        "Журнал подій",
        ["День", "Репетитор", "Учень", "Подія", "Канал", "Сума, грн", "Примітка"],
        [
            [
                r["day"],
                r["tutor"],
                r["student"],
                EVENTS[r["event"]],
                r["channel"],
                r["amount"] or "",
                r["detail"],
            ]
            for r in rows
        ],
    )
    sheet(
        wb,
        "Щоденник",
        [
            "День",
            "Нагадувань",
            "Боргів з вікном 3 дні",
            "Погашено за ≤3 дні",
            "Частка, %",
            "Негативних",
            "Помилкових",
        ],
        [
            [
                d["day"],
                d["reminders"],
                d["episodes"],
                d["paid3"],
                d["share3"],
                d["negative"],
                d["false"],
            ]
            for d in m["diary"]
        ],
    )
    sheet(
        wb,
        "Репетитори",
        [
            "Репетитор",
            "Учнів",
            "Занять",
            "Нагадувань",
            "Боргів",
            "Погашено ≤3 дні",
            "Звітів, тиж. 1",
            "Звітів, тиж. 2",
            "Помилкових",
            "Негативних",
        ],
        [
            [
                t,
                v["students"],
                v["lessons"],
                v["reminders"],
                v["episodes"],
                v["paid3"],
                v["report_w1"],
                v["report_w2"],
                v["false"],
                v["negative"],
            ]
            for t, v in m["by_tutor"].items()
        ],
    )
    wb.save(OUT)
    print(f"-> {OUT.name}")


if __name__ == "__main__":
    main()
